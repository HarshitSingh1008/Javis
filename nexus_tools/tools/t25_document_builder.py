"""
NEXUS AI v4.0 — Tool 25: Document compilation (PowerPoint).
Hardware: Intel i3 7th Gen · 12GB RAM · Ollama + Groq API

Compiles PowerPoint (.pptx) presentations from structured slide data.
Matches the original JARVIS document_builder tool.
"""

import json
import logging
import os
import time
from typing import List, Dict, Any
from langchain_core.tools import tool

logger = logging.getLogger("nexus.tool.document_builder")


@tool
def document_builder(
    filepath: str,
    title: str,
    slide_data: List[Dict[str, str]],
) -> str:
    """
    Compiles a PowerPoint (.pptx) file natively.
    
    Use this tool when: The user asks to create a presentation, slideshow, or PowerPoint document.
    
    Args:
        filepath: Absolute path to save the .pptx file.
        title: Main title of the presentation.
        slide_data: List of dictionaries representing slides.
                    Format: [{"title": "Slide Title", "content": "Bullet 1\nBullet 2\nBullet 3"}]
    
    Returns:
        JSON string with keys:
          - success (bool): Whether the compilation succeeded.
          - result (str): Save confirmation with path.
          - error (str or null): Error message if failed.
    
    Examples:
        >>> document_builder("C:/presentation.pptx", "My Presentation", [{"title": "Intro", "content": "Point 1\nPoint 2"}])
    """
    start = time.perf_counter()
    
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        
        prs = Presentation()
        
        # Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title_shape = slide.shapes.title
        subtitle = slide.placeholders[1]
        title_shape.text = title
        subtitle.text = "Generated autonomously by NEXUS AI"
        
        # Content Slides
        bullet_slide_layout = prs.slide_layouts[1]
        for data in slide_data:
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]
            
            title_shape.text = data.get("title", "Untitled Slide")
            tf = body_shape.text_frame
            tf.text = data.get("content", "")
            
            # Format text
            for paragraph in tf.paragraphs:
                paragraph.font.size = Pt(14)
        
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        prs.save(filepath)
        
        return json.dumps({
            "success": True,
            "result": f"PowerPoint presentation successfully compiled and saved to {filepath}",
            "error": None,
        })
    
    except Exception as e:
        logger.error(f"document_builder error: {e}", exc_info=True)
        return json.dumps({
            "success": False, "result": None,
            "error": f"{type(e).__name__}: {e}"
        })